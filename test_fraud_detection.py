import io
import unittest
import zipfile

import pymupdf
from docx import Document

from backend.fraud_detection import analyze_document


def _docx_bytes(include_hidden_text=False):
    document = Document()
    document.add_paragraph("Application document")
    if include_hidden_text:
        run = document.add_paragraph().add_run("Hidden review text")
        run.font.hidden = True
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


class FraudDetectionTests(unittest.TestCase):
    def test_clean_pdf_returns_pdf_indicators_shape(self):
        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "A normal application document.")
        result = analyze_document(pdf.tobytes(), ".pdf", "sample.pdf")
        self.assertEqual(result["method"], "pdf_metadata_and_structure_heuristics")
        self.assertIsInstance(result["indicators"], list)

    def test_pdf_javascript_is_reported_without_execution(self):
        pdf = pymupdf.open()
        pdf.new_page()
        pdf.xref_set_key(
            pdf.pdf_catalog(),
            "OpenAction",
            "<< /S /JavaScript /JS (app.alert\\('test'\\)) >>",
        )
        result = analyze_document(pdf.tobytes(), ".pdf", "sample.pdf")
        signals = {item["signal"] for item in result["indicators"]}
        self.assertIn("pdf_document_action", signals)
        self.assertIn("pdf_javascript", signals)
        self.assertGreaterEqual(result["risk_score"], 70)

    def test_docx_hidden_text_is_reported(self):
        result = analyze_document(_docx_bytes(include_hidden_text=True), ".docx", "sample.docx")
        signals = {item["signal"] for item in result["indicators"]}
        self.assertIn("docx_hidden_text", signals)
        self.assertEqual(result["method"], "docx_metadata_and_structure_heuristics")

    def test_docx_external_relationship_is_reported(self):
        source = _docx_bytes()
        updated = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(source)) as original, zipfile.ZipFile(updated, "w") as output:
            for item in original.infolist():
                contents = original.read(item.filename)
                if item.filename == "word/_rels/document.xml.rels":
                    contents = contents.replace(
                        b"</Relationships>",
                        b'<Relationship Id="rId999" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://example.invalid/" TargetMode="External"/></Relationships>',
                    )
                output.writestr(item, contents)
        result = analyze_document(updated.getvalue(), ".docx", "sample.docx")
        signals = {item["signal"] for item in result["indicators"]}
        self.assertIn("docx_external_relationships", signals)

    def test_malformed_pdf_is_rejected(self):
        with self.assertRaises(ValueError):
            analyze_document(b"%PDF-invalid", ".pdf", "sample.pdf")


if __name__ == "__main__":
    unittest.main()
