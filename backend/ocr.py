"""Document text extraction with OpenCV preprocessing and Tesseract OCR."""

from __future__ import annotations

import io
import os
import re
import shutil
import zipfile
from typing import Any

import cv2
import pymupdf as fitz
import numpy as np
import pytesseract
from docx import Document
from PIL import Image, ImageOps


class OCRFailure(Exception):
    """A readable extraction error suitable for an API response."""


def extract_text(content: bytes, suffix: str) -> dict[str, Any]:
    if suffix == ".docx":
        return _extract_docx(content)
    if suffix == ".pdf":
        return _extract_pdf(content)
    return _extract_image(content)


def _extract_docx(content: bytes) -> dict[str, Any]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if len(archive.infolist()) > 500 or sum(item.file_size for item in archive.infolist()) > 30_000_000:
                raise OCRFailure("The DOCX document contains too much expanded content.")
    except zipfile.BadZipFile as exc:
        raise OCRFailure("The DOCX document is not a valid Office file.") from exc
    try:
        document = Document(io.BytesIO(content))
    except Exception as exc:
        raise OCRFailure("The DOCX document could not be opened.") from exc
    text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
    for table in document.tables:
        text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    if not text.strip():
        raise OCRFailure("No readable text was found in the DOCX document.")
    return {"text": text, "pages": 1, "language": "eng", "quality": _quality(text), "warnings": []}


def _extract_pdf(content: bytes) -> dict[str, Any]:
    try:
        pdf = fitz.open(stream=content, filetype="pdf")
    except Exception as exc:
        raise OCRFailure("The PDF could not be opened.") from exc
    if pdf.is_encrypted:
        raise OCRFailure("Password-protected PDFs are not supported.")
    if pdf.page_count > 30:
        raise OCRFailure("PDFs are limited to 30 pages.")
    pieces: list[str] = []
    for page in pdf:
        text = page.get_text("text").strip()
        if len(text) < 30:
            scale = min(1.7, 3000 / max(page.rect.width, page.rect.height))
            pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
            image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
            text = _ocr_image(image)
        pieces.append(text)
    text = "\n".join(piece for piece in pieces if piece.strip())
    if not text.strip():
        raise OCRFailure("No readable text was found. Try a clearer scan or a supported language.")
    return {
        "text": text,
        "pages": pdf.page_count,
        "language": os.getenv("TESSERACT_LANG", "eng"),
        "quality": _quality(text),
        "warnings": _quality_warnings(text),
    }


def _extract_image(content: bytes) -> dict[str, Any]:
    try:
        image = Image.open(io.BytesIO(content))
        if image.width * image.height > 30_000_000:
            raise OCRFailure("Images larger than 30 megapixels are not supported.")
        image = image.convert("RGB")
    except OCRFailure:
        raise
    except Exception as exc:
        raise OCRFailure("The image could not be decoded.") from exc
    text = _ocr_image(image)
    if not text.strip():
        raise OCRFailure("No readable text was found. Try a higher-resolution scan.")
    return {
        "text": text,
        "pages": 1,
        "language": os.getenv("TESSERACT_LANG", "eng"),
        "quality": _quality(text),
        "warnings": _quality_warnings(text),
    }


def _ocr_image(image: Image.Image) -> str:
    if shutil.which(os.getenv("TESSERACT_CMD", "tesseract")) is None:
        configured = os.getenv("TESSERACT_CMD")
        if configured:
            pytesseract.pytesseract.tesseract_cmd = configured
        elif os.name == "nt":
            default_path = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
            if os.path.exists(default_path):
                pytesseract.pytesseract.tesseract_cmd = default_path
            else:
                raise OCRFailure("Tesseract is not installed. Install it and add it to PATH.")
        else:
            raise OCRFailure("Tesseract is not installed. Install it and add it to PATH.")
    # OpenCV grayscale + adaptive threshold helps with uneven lighting on scans.
    cv_image = cv2.cvtColor(np.array(ImageOps.exif_transpose(image)), cv2.COLOR_RGB2GRAY)
    if max(cv_image.shape) < 1400:
        scale = 1400 / max(cv_image.shape)
        cv_image = cv2.resize(cv_image, None, fx=scale, fy=scale, interpolation=cv2.INTER_CUBIC)
    pad_y = max(0, 31 - cv_image.shape[0])
    pad_x = max(0, 31 - cv_image.shape[1])
    if pad_y or pad_x:
        cv_image = cv2.copyMakeBorder(cv_image, 0, pad_y, 0, pad_x, cv2.BORDER_CONSTANT, value=255)
    processed = cv2.adaptiveThreshold(
        cv_image, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C, cv2.THRESH_BINARY, 31, 11
    )
    try:
        return pytesseract.image_to_string(
            processed, lang=os.getenv("TESSERACT_LANG", "eng"), config="--psm 6"
        ).strip()
    except pytesseract.TesseractNotFoundError as exc:
        raise OCRFailure("Tesseract is not installed or TESSERACT_CMD is misconfigured.") from exc
    except pytesseract.TesseractError as exc:
        raise OCRFailure(f"OCR failed: {exc}") from exc


def _quality(text: str) -> int:
    if not text:
        return 0
    printable = sum(character.isprintable() for character in text)
    tokens = re.findall(r"\b[\w-]+\b", text)
    return max(0, min(100, round(100 * printable / len(text)) - (25 if len(tokens) < 5 else 0)))


def _quality_warnings(text: str) -> list[str]:
    warnings = []
    if len(text.strip()) < 40:
        warnings.append("Little text was extracted; verify scan quality manually.")
    if len(re.findall(r"[�□]", text)) > 3:
        warnings.append("Several characters may be unreadable or unsupported.")
    return warnings